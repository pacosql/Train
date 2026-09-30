"""Genera mates10/excel/mates10.xlsx: una hoja por tabla principal y tablas dinámicas NATIVAS de Excel
(creadas con LibreOffice Calc por UNO y exportadas a OOXML), más la lista de decisiones.

    python3 mates10/tools/excel.py
"""
import os
import re
import subprocess
import sys
import tempfile
import time

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from db import q  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SALIDA = os.path.join(ROOT, "excel", "mates10.xlsx")

HOJAS = {
    "habilidades": """select h.codigo, b.nombre bloque, h.familia, h.nivel_familia, h.curso_ref curso, h.nombre, h.nombre_padres, h.definicion,
        h.ejemplo_frontera, h.ejemplo_fuera, h.importancia, h.carga_alumno, case when h.generador is null then 'IA' else 'código' end generador,
        h.estado_revision, f.codigo fuente, h.fuente_ref,
        (select count(*) from mates10_error_tipico e where e.habilidad_id = h.id) errores_tipicos,
        (select count(*) from mates10_habilidad_prerrequisito p where p.habilidad_id = h.id) prerrequisitos,
        (select count(*) from mates10_ejercicio_habilidad eh where eh.habilidad_id = h.id and eh.rol = 'principal') ejercicios
        from mates10_habilidad h join mates10_bloque b on b.id = h.bloque_id join mates10_fuente f on f.id = h.fuente_id order by h.codigo""",
    "errores_tipicos": """select e.codigo, h.codigo habilidad, e.nombre, e.descripcion, e.explicacion_nino, e.indicacion_padres, f.codigo fuente,
        (select count(*) from mates10_ejercicio_error ee where ee.error_tipico_id = e.id) ejercicios_que_lo_detectan
        from mates10_error_tipico e join mates10_habilidad h on h.id = e.habilidad_id left join mates10_fuente f on f.id = e.fuente_id order by e.codigo""",
    "prerrequisitos": """select h.codigo habilidad, p.codigo prerrequisito, hp.peso from mates10_habilidad_prerrequisito hp
        join mates10_habilidad h on h.id = hp.habilidad_id join mates10_habilidad p on p.id = hp.prerrequisito_id order by 1, 2""",
    "mapa": """select s.codigo sistema, s.nombre comunidad, c.codigo curso, m.trimestre, m.mes, m.orden, h.codigo habilidad, b.nombre bloque, h.familia,
        m.estado, m.confianza, coalesce(m.editorial, '') editorial, f.codigo fuente, m.fuente_ref, m.nota
        from mates10_mapa_curricular m join mates10_sistema_educativo s on s.id = m.sistema_id join mates10_curso c on c.id = m.curso_id
        join mates10_habilidad h on h.id = m.habilidad_id join mates10_bloque b on b.id = h.bloque_id join mates10_fuente f on f.id = m.fuente_id
        order by s.codigo, c.orden, m.trimestre, m.mes nulls first, m.orden""",
    "ejercicios": """select h.codigo habilidad, h.curso_ref curso, b.nombre bloque, e.formato, e.dificultad, e.enunciado, e.respuesta->>0 respuesta,
        e.distractores::text distractores, case when e.generado_por = 'algoritmo' then 'código' else 'IA' end generador,
        case when e.rechazado then 'rechazado' when e.revisado_por is not null then 'revisado (' || e.revision_tipo || ')'
             when e.publicado then 'provisional' else 'pendiente' end estado_revision, e.publicado
        from mates10_ejercicio e join mates10_ejercicio_habilidad eh on eh.ejercicio_id = e.id and eh.rol = 'principal'
        join mates10_habilidad h on h.id = eh.habilidad_id join mates10_bloque b on b.id = h.bloque_id order by h.codigo, e.dificultad""",
    "metodos": """select m.codigo, m.nombre, m.escuela, m.representacion, coalesce(m.familia, '') familia, m.descripcion, m.cuando_usar, m.edad_desde, f.codigo fuente
        from mates10_metodo m join mates10_fuente f on f.id = m.fuente_id order by m.familia, m.codigo""",
    "metodos_por_familia": """select h.familia, m.codigo metodo, m.escuela, hm.rol, count(*) habilidades from mates10_habilidad_metodo hm
        join mates10_habilidad h on h.id = hm.habilidad_id join mates10_metodo m on m.id = hm.metodo_id group by 1, 2, 3, 4 order by 1, 4, 2""",
    "tecnicas": """select t.codigo, t.nombre, coalesce(t.regla->>'_tipo', 'tecnica_practica') tipo, t.disparador, t.ambito, t.descripcion, t.regla::text regla, f.codigo fuente
        from mates10_tecnica_practica t join mates10_fuente f on f.id = t.fuente_id order by 3, 1""",
    "tecnicas_por_familia": """select h.familia, t.codigo tecnica, min(ht.prioridad) prioridad, count(*) habilidades from mates10_habilidad_tecnica ht
        join mates10_habilidad h on h.id = ht.habilidad_id join mates10_tecnica_practica t on t.id = ht.tecnica_id group by 1, 2 order by 1, 3""",
    "mecanismos_enganche": """select m.codigo, m.nombre, m.evidencia, m.riesgo_menores, m.implementable, m.estado, m.descripcion, m.uso_en_mates10, m.origen,
        m.edad_desde, m.edad_hasta, f.codigo fuente from mates10_mecanismo_enganche m join mates10_fuente f on f.id = m.fuente_id order by m.estado, m.codigo""",
    "fuentes": """select f.codigo, f.tipo, f.titulo, coalesce(f.editorial, '') editorial, f.anio, f.origen, f.confianza, array_to_string(f.modulos, ', ') modulos,
        coalesce(s.codigo, '') sistema, f.url_original, f.capturado_en::date capturado, f.notas,
        (select count(*) from mates10_fuente_archivo a where a.fuente_id = f.id) archivos
        from mates10_fuente f left join mates10_sistema_educativo s on s.id = f.sistema_id order by f.tipo, f.codigo""",
    "fuentes_por_modulo": """select f.codigo, f.tipo, f.confianza, unnest(case when f.modulos = '{}' then array['sin módulo'] else f.modulos end) modulo
        from mates10_fuente f order by 4, 2""",
    "colegios": """select s.codigo sistema, c.codigo_oficial, c.nombre, c.municipio, c.tipo, array_to_string(c.ensenanzas, ', ') ensenanzas
        from mates10_colegio c join mates10_sistema_educativo s on s.id = c.sistema_id order by c.municipio, c.nombre""",
    "sistemas_y_cursos": """select s.codigo sistema, s.nombre, c.codigo curso, c.nombre_local, c.edad_tipica, c.etapa
        from mates10_curso c join mates10_sistema_educativo s on s.id = c.sistema_id order by s.codigo, c.orden""",
}

# Tablas dinámicas: (nombre de hoja, hoja origen, filas, columnas, campo de datos, función)
PIVOTS = [
    ("TD habilidades curso-bloque", "habilidades", ["curso"], ["bloque"], "codigo", "COUNT"),
    ("TD ejercicios habilidad-estado", "ejercicios", ["curso", "habilidad"], ["estado_revision"], "enunciado", "COUNT"),
    ("TD ejercicios formato-generador", "ejercicios", ["bloque", "formato"], ["generador"], "enunciado", "COUNT"),
    ("TD mapa comunidad-curso", "mapa", ["comunidad"], ["curso"], "habilidad", "COUNT"),
    ("TD mapa confianza", "mapa", ["comunidad"], ["confianza"], "habilidad", "COUNT"),
    ("TD fuentes tipo-módulo", "fuentes_por_modulo", ["tipo"], ["modulo"], "codigo", "COUNT"),
    ("TD métodos por familia", "metodos_por_familia", ["familia"], ["escuela"], "metodo", "COUNT"),
    ("TD técnicas por familia", "tecnicas_por_familia", ["familia"], ["tecnica"], "prioridad", "MIN"),
    ("TD enganche evidencia-riesgo", "mecanismos_enganche", ["evidencia"], ["riesgo_menores"], "codigo", "COUNT"),
]


def decisiones():
    txt = open(os.path.join(ROOT, "DECISIONES.md"), encoding="utf-8").read()
    filas, seccion = [], ""
    for linea in txt.splitlines():
        if linea.startswith("#"):
            seccion = linea.strip("# ").strip()
            continue
        m = re.match(r"^- \*\*(D-\d+) · ([^·]+) · ([^*]+)\.\*\* (.*)$", linea)
        if m:
            filas.append({"codigo": m.group(1), "fecha": m.group(2).strip(), "area": m.group(3).strip(), "seccion": seccion, "decision": m.group(4)})
        elif linea.startswith("- **") and seccion:
            filas.append({"codigo": "", "fecha": "", "area": "", "seccion": seccion, "decision": re.sub(r"\*\*", "", linea[2:])})
    return filas


def escribir_datos(ruta):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    info = wb.create_sheet("LEEME")
    info.append(["Mates10 — banco de contenidos. Generado automáticamente desde Supabase."])
    info.append(["Hojas 'TD …': tablas dinámicas de Excel (clic dentro → panel de campos para cambiar filas, columnas y filtros)."])
    info.append(["Resto: una hoja por tabla principal. Las decisiones están en 'decisiones'."])
    info.append([f"Fecha: {time.strftime('%Y-%m-%d %H:%M')}"])
    dims = {}
    for nombre, sql in list(HOJAS.items()) + [("decisiones", None)]:
        filas = q(sql) if sql else decisiones()
        ws = wb.create_sheet(nombre[:31])
        cols = list(filas[0].keys()) if filas else ["(vacío)"]
        ws.append(cols)
        for r in filas:
            ws.append([("" if r[c] is None else r[c]) if not isinstance(r[c], (dict, list)) else str(r[c]) for c in cols])
        for i, c in enumerate(cols, 1):
            ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
            ws.cell(1, i).fill = PatternFill("solid", fgColor="D97757")
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(c) + 2))
        ws.freeze_panes = "A2"
        if filas:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(filas) + 1}"
        dims[nombre] = (cols, len(filas))
        print(f"  {nombre}: {len(filas)} filas")
    wb.save(ruta)
    return dims


def anadir_pivots(entrada, salida, dims):
    import uno
    from com.sun.star.beans import PropertyValue
    from com.sun.star.table import CellAddress, CellRangeAddress
    from com.sun.star.sheet.DataPilotFieldOrientation import ROW, COLUMN, DATA

    puerto = 2100 + os.getpid() % 500
    perfil = tempfile.mkdtemp(prefix="lo-")
    proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore", f"-env:UserInstallation=file://{perfil}",
                             f"--accept=socket,host=localhost,port={puerto};urp;"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        ctx = None
        for _ in range(80):
            try:
                local = uno.getComponentContext()
                res = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
                ctx = res.resolve(f"uno:socket,host=localhost,port={puerto};urp;StarOffice.ComponentContext")
                break
            except Exception:  # noqa: BLE001
                time.sleep(0.5)
        desk = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)

        def pv(n, v):
            x = PropertyValue()
            x.Name, x.Value = n, v
            return x
        doc = desk.loadComponentFromURL(uno.systemPathToFileUrl(entrada), "_blank", 0, (pv("Hidden", True),))
        hojas = doc.Sheets
        pos = 1
        for nombre, origen, filas, columnas, dato, func in PIVOTS:
            cols, n = dims[origen]
            if n == 0:
                continue
            src = hojas.getByName(origen)
            hojas.insertNewByName(nombre[:31], pos)
            destino = hojas.getByName(nombre[:31])
            pos += 1
            dp = destino.getDataPilotTables()
            d = dp.createDataPilotDescriptor()
            rng = CellRangeAddress()
            rng.Sheet = src.getRangeAddress().Sheet
            rng.StartColumn, rng.StartRow, rng.EndColumn, rng.EndRow = 0, 0, len(cols) - 1, n
            d.setSourceRange(rng)
            campos = d.getDataPilotFields()
            idx = {campos.getByIndex(i).Name: i for i in range(campos.getCount())}
            for f in filas:
                campos.getByIndex(idx[f]).Orientation = ROW
            for f in columnas:
                campos.getByIndex(idx[f]).Orientation = COLUMN
            df = campos.getByIndex(idx[dato])
            df.Orientation = DATA
            df.Function = uno.Enum("com.sun.star.sheet.GeneralFunction", func)
            a = CellAddress()
            a.Sheet = destino.getRangeAddress().Sheet
            a.Column, a.Row = 0, 2
            destino.getCellByPosition(0, 0).setString(f"{nombre} · origen: hoja «{origen}»")
            dp.insertNewByName("td" + str(pos), a, d)
        doc.storeToURL(uno.systemPathToFileUrl(salida), (pv("FilterName", "Calc MS Excel 2007 XML"),))
        doc.close(True)
    finally:
        proc.terminate()


if __name__ == "__main__":
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    tmp = os.path.join(tempfile.mkdtemp(), "datos.xlsx")
    print("Datos:")
    dims = escribir_datos(tmp)
    print("Tablas dinámicas…")
    anadir_pivots(tmp, SALIDA, dims)
    print("Escrito", SALIDA, os.path.getsize(SALIDA), "bytes")
